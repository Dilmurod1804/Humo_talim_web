"""
HUMO O'quv markazi — Staff, Face ID, Salary, Bonus, Advance & Audit Views (face_id_views.py)
O'qituvchilar, menejerlar, Face ID biometrik davomat, oylik maosh, bonuslar, avanslar va audit.
"""
import base64
import json
import os
import random
import string
import uuid
from datetime import date, datetime
from decimal import Decimal

import openpyxl
from django.contrib import messages
from django.contrib.auth import login
from django.contrib.auth.decorators import login_required
from django.core.files.base import ContentFile
from django.db.models import Count, Q, Sum
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.views.decorators.csrf import csrf_exempt

from .auth_views import admin_only_required, admin_required, is_admin_user
from .face_service import (
    calculate_euclidean_distance,
    calculate_haversine_distance,
    calculate_staff_salary_summary,
    compare_embeddings,
    compare_two_faces,
    decode_base64_image,
    extract_face_embedding,
    get_or_create_face_embedding,
    get_user_face_data,
    get_user_staff_profile,
    match_staff_by_face,
    parse_descriptor,
    process_face_check_in,
    save_staff_descriptor,
    sync_all_face_embeddings,
    sync_staff_face_embedding,
)
from .models import (
    AdvanceRequest, AuditLog, CustomUser, FaceIDLog, FaceIDSettings,
    Group, Manager, StaffBonus, Teacher,
)
from .payment_views import _is_ajax, normalize_month_str
from .services import optimize_student_photo
from .validators import contains_cyrillic


# ─────────────────────────────────────────────────────────────
# 1. YORDAMCHI FUNKSIYALAR (HELPERS)
# ─────────────────────────────────────────────────────────────

def _safe_decimal(value, default='0') -> Decimal:
    """Har qanday string qiymatni xavfsiz ravishda Decimal ga o'giradi."""
    try:
        cleaned = str(value).replace(',', '.').strip()
        if not cleaned or cleaned in ('.', '-'):
            return Decimal(default)
        return Decimal(cleaned)
    except Exception:
        return Decimal(default)


def _generate_staff_credentials(prefix: str, first_name: str) -> tuple[str, str]:
    login_str = f"{prefix}_{first_name.lower()[:3]}_{random.randint(100, 999)}"
    pwd_str = ''.join(random.choices(string.ascii_letters + string.digits, k=8))
    return login_str, pwd_str


def _check_staff_face_detected(staff_or_user, has_cv: bool = True, fs_module=None) -> bool:
    """
    Xodimning (`user.teacher_profile` yoki `user.manager_profile`) Face ID biometrik shabloni
    (`face_embedding`) mavjudligini tekshiradi.
    Agar xodimda rasm (`photo`) biriktirilgan bo'lsa-yu, lekin Face ID bazasida uning
    shabloni (`face_embedding`) bo'lmasa, uni avtomatik generatsiya qilib bazaga bog'laydi.
    """
    staff_obj, photo_field, _ = get_user_face_data(staff_or_user)
    if not staff_obj:
        return False

    # 1. Agar bazada tayyor face_embedding shabloni bor bo'lsa -> True
    if getattr(staff_obj, 'has_face_embedding', False):
        return True

    # 2. Agar rasmi bor bo'lsa, lekin face_embedding bo'sh bo'lsa -> avtomatik bog'laymiz
    if photo_field:
        emb = get_or_create_face_embedding(staff_obj)
        if emb is not None:
            return True
        if sync_staff_face_embedding(staff_obj):
            return True

    return False


def _serialize_staff_face_status(staff_qs, role: str, role_title: str, has_cv: bool, fs_module) -> list[dict]:
    """Menejer yoki O'qituvchilar ro'yxatidan Face ID holati kartalarini shakllantiradi."""
    items = []
    for s in staff_qs:
        has_photo = bool(s.photo)
        face_detected = _check_staff_face_detected(s, has_cv, fs_module) if (has_photo or s.has_face_embedding) else False
        items.append({
            'id': s.id,
            'role': role,
            'role_title': role_title,
            'full_name': f"{s.first_name} {s.last_name}".strip() or s.user.username,
            'username': s.user.username,
            'phone': s.phone or '—',
            'has_photo': has_photo,
            'face_detected': face_detected,
            'photo_url': s.photo.url if has_photo else '',
            'work_time': f"{s.work_start_time.strftime('%H:%M')} - {s.work_end_time.strftime('%H:%M')}" if s.work_start_time else '—',
            'is_active': s.is_active,
        })
    return items


# ─────────────────────────────────────────────────────────────
# 2. O'QITUVCHILAR VA MENAJERLAR BOSHQARUVI (CRUD)
# ─────────────────────────────────────────────────────────────

@admin_only_required
def admin_teacher_create(request):
    settings_obj = FaceIDSettings.get_settings()
    default_ctx = {
        'default_fine_rate': settings_obj.default_late_fine_rate,
        'default_start_time': settings_obj.default_work_start_time,
        'default_end_time': settings_obj.default_work_end_time,
    }
    if request.method == 'POST':
        first_name = request.POST.get('first_name', '').strip()
        last_name = request.POST.get('last_name', '').strip()
        phone = request.POST.get('phone', '').strip()
        monthly_salary = request.POST.get('monthly_salary', '0').strip()
        late_fine_rate = request.POST.get('late_fine_rate', str(settings_obj.default_late_fine_rate)).strip()
        work_start_time = request.POST.get('work_start_time', str(settings_obj.default_work_start_time)).strip()
        work_end_time = request.POST.get('work_end_time', str(settings_obj.default_work_end_time)).strip()
        photo_file = request.FILES.get('photo')

        if contains_cyrillic(first_name) or contains_cyrillic(last_name):
            messages.error(request, "Ism va familiyada faqat lotin alifbosidan foydalaning (Kirill harflari taqiqlangan)!")
            return render(request, 'admin_teacher_create.html', {
                **default_ctx,
                'first_name_val': first_name,
                'last_name_val': last_name,
                'phone_val': phone,
            })

        login_str, pwd_str = _generate_staff_credentials('t', first_name)
        user = CustomUser.objects.create_user(username=login_str, password=pwd_str, role='teacher')
        opt_photo, _ = optimize_student_photo(photo_file) if photo_file else (None, None)

        salary_val = _safe_decimal(monthly_salary, '0') if is_admin_user(request) else Decimal('0')
        fine_rate_val = (
            _safe_decimal(late_fine_rate, str(settings_obj.default_late_fine_rate))
            if is_admin_user(request) else settings_obj.default_late_fine_rate
        )
        teacher = Teacher.objects.create(
            user=user,
            first_name=first_name,
            last_name=last_name,
            phone=phone,
            generated_login=login_str,
            generated_password=pwd_str,
            monthly_salary=salary_val,
            late_fine_rate=fine_rate_val,
            work_start_time=work_start_time or '09:00',
            work_end_time=work_end_time or '18:00',
            photo=opt_photo,
        )
        if teacher.photo:
            sync_staff_face_embedding(teacher)
        messages.success(request, f"O'qituvchi '{first_name} {last_name}' muvaffaqiyatli yaratildi (Face ID fotosi biriktirildi).")
        return redirect('admin_dashboard')
    return render(request, 'admin_teacher_create.html', default_ctx)


@admin_only_required
def admin_teacher_edit(request, id):
    teacher = get_object_or_404(Teacher, id=id)
    if request.method == 'POST':
        first_name = request.POST.get('first_name', '').strip()
        last_name = request.POST.get('last_name', '').strip()
        if contains_cyrillic(first_name) or contains_cyrillic(last_name):
            messages.error(request, "Ism va familiyada faqat lotin alifbosidan foydalaning!")
            return render(request, 'admin_teacher_edit.html', {'teacher': teacher})

        teacher.first_name = first_name
        teacher.last_name = last_name
        teacher.phone = request.POST.get('phone', '').strip()
        if is_admin_user(request):
            teacher.monthly_salary = _safe_decimal(request.POST.get('monthly_salary', '0'), str(teacher.monthly_salary))
            teacher.late_fine_rate = _safe_decimal(request.POST.get('late_fine_rate', '2.0'), str(teacher.late_fine_rate))
        teacher.work_start_time = request.POST.get('work_start_time', '09:00').strip() or '09:00'
        teacher.work_end_time = request.POST.get('work_end_time', '18:00').strip() or '18:00'
        teacher.can_edit_students = request.POST.get('can_edit_students') == '1'

        photo_file = request.FILES.get('photo')
        if photo_file:
            opt_photo, _ = optimize_student_photo(photo_file)
            if opt_photo:
                teacher.photo = opt_photo

        new_password = request.POST.get('new_password', '').strip()
        if new_password:
            teacher.user.set_password(new_password)
            teacher.user.save()
            teacher.generated_password = new_password

        teacher.save()
        if photo_file or (teacher.photo and not teacher.has_face_embedding):
            sync_staff_face_embedding(teacher)
        messages.success(request, f"O'qituvchi '{teacher.first_name} {teacher.last_name}' ma'lumotlari muvaffaqiyatli saqlandi.")
        return redirect('admin_dashboard')

    return render(request, 'admin_teacher_edit.html', {'teacher': teacher})


@admin_only_required
def admin_teacher_delete(request, id):
    teacher = get_object_or_404(Teacher, id=id)
    if request.method == 'POST':
        user = teacher.user
        teacher.delete()
        if user:
            user.delete()
        messages.success(request, "O'qituvchi o'chirildi.")
    return redirect('admin_dashboard')


@admin_only_required
def admin_teacher_toggle_edit(request, id):
    """Admin o'qituvchiga o'quvchini tahrirlash ruxsatini yoqadi/o'chiradi."""
    teacher = get_object_or_404(Teacher, id=id)
    teacher.can_edit_students = not bool(teacher.can_edit_students)
    teacher.save(update_fields=['can_edit_students'])
    status_text = "yoqildi ✅" if teacher.can_edit_students else "o'chirildi ❌"
    messages.success(request, f"{teacher.first_name} {teacher.last_name} uchun tahrirlash ruxsati {status_text}")
    return redirect('admin_dashboard')


@admin_only_required
def admin_teacher_toggle_all_edit(request):
    """Barcha o'qituvchilar uchun tahrirlash ruxsatini bir yo'la yoqish/o'chirish."""
    if request.method == 'POST':
        new_val = request.POST.get('action', 'off') == 'on'
        count = Teacher.objects.update(can_edit_students=new_val)
        state = "YOQILDI ✅" if new_val else "O'CHIRILDI ❌"
        messages.success(request, f"Barcha {count} ta o'qituvchi uchun tahrirlash ruxsati {state}")
    return redirect('admin_dashboard')


@admin_only_required
def admin_manager_create(request):
    settings_obj = FaceIDSettings.get_settings()
    if request.method == 'POST':
        first_name = request.POST.get('first_name', '').strip()
        last_name = request.POST.get('last_name', '').strip()
        phone = request.POST.get('phone', '').strip()
        monthly_salary = request.POST.get('monthly_salary', '0').strip()
        late_fine_rate = request.POST.get('late_fine_rate', '2.0').strip()
        work_start_time = request.POST.get('work_start_time', '09:00').strip()
        work_end_time = request.POST.get('work_end_time', '18:00').strip()
        photo_file = request.FILES.get('photo')

        if contains_cyrillic(first_name) or contains_cyrillic(last_name):
            messages.error(request, "Ism va familiyada faqat lotin alifbosidan foydalaning!")
            return render(request, 'admin_manager_create.html', request.POST)

        login_str, pwd_str = _generate_staff_credentials('m', first_name)
        user = CustomUser.objects.create_user(username=login_str, password=pwd_str, role='manager')
        opt_photo, _ = optimize_student_photo(photo_file) if photo_file else (None, None)

        salary_val = _safe_decimal(monthly_salary, '0') if is_admin_user(request) else Decimal('0')
        fine_rate_val = (
            _safe_decimal(late_fine_rate, str(settings_obj.default_late_fine_rate))
            if is_admin_user(request) else settings_obj.default_late_fine_rate
        )
        manager = Manager.objects.create(
            user=user,
            first_name=first_name,
            last_name=last_name,
            phone=phone,
            generated_login=login_str,
            generated_password=pwd_str,
            monthly_salary=salary_val,
            late_fine_rate=fine_rate_val,
            work_start_time=work_start_time or '09:00',
            work_end_time=work_end_time or '18:00',
            photo=opt_photo,
        )
        if manager.photo:
            sync_staff_face_embedding(manager)
        messages.success(request, f"Menejer '{first_name} {last_name}' muvaffaqiyatli qo'shildi.")
        return redirect('admin_dashboard')

    return render(request, 'admin_manager_create.html', {
        'default_fine_rate': settings_obj.default_late_fine_rate,
        'default_start_time': settings_obj.default_work_start_time,
        'default_end_time': settings_obj.default_work_end_time,
    })


@admin_only_required
def admin_manager_edit(request, id):
    manager = get_object_or_404(Manager, id=id)
    if request.method == 'POST':
        first_name = request.POST.get('first_name', '').strip()
        last_name = request.POST.get('last_name', '').strip()
        if contains_cyrillic(first_name) or contains_cyrillic(last_name):
            messages.error(request, "Ism va familiyada faqat lotin alifbosidan foydalaning!")
            return render(request, 'admin_manager_edit.html', {'manager': manager})

        manager.first_name = first_name
        manager.last_name = last_name
        manager.phone = request.POST.get('phone', '').strip()
        if is_admin_user(request):
            manager.monthly_salary = _safe_decimal(request.POST.get('monthly_salary', '0'), str(manager.monthly_salary))
            manager.late_fine_rate = _safe_decimal(request.POST.get('late_fine_rate', '2.0'), str(manager.late_fine_rate))
        manager.work_start_time = request.POST.get('work_start_time', '09:00').strip() or '09:00'
        manager.work_end_time = request.POST.get('work_end_time', '18:00').strip() or '18:00'

        photo_file = request.FILES.get('photo')
        if photo_file:
            opt_photo, _ = optimize_student_photo(photo_file)
            if opt_photo:
                manager.photo = opt_photo

        new_password = request.POST.get('new_password', '').strip()
        if new_password:
            manager.user.set_password(new_password)
            manager.user.save()
            manager.generated_password = new_password

        manager.save()
        if photo_file or (manager.photo and not manager.has_face_embedding):
            sync_staff_face_embedding(manager)
        messages.success(request, f"Menejer '{manager.first_name} {manager.last_name}' ma'lumotlari yangilandi.")
        return redirect('admin_dashboard')

    return render(request, 'admin_manager_edit.html', {'manager': manager})


@admin_only_required
def admin_manager_delete(request, id):
    manager = get_object_or_404(Manager, id=id)
    if request.method == 'POST':
        user = manager.user
        name = f"{manager.first_name} {manager.last_name}"
        manager.delete()
        if user:
            user.delete()
        messages.success(request, f"Menejer '{name}' tizimdan o'chirildi.")
    return redirect('admin_dashboard')


# ─────────────────────────────────────────────────────────────
# 3. FACE ID JURNALI, SINXRONIZATSIYA VA SOZLAMALAR
# ─────────────────────────────────────────────────────────────

@admin_only_required
def admin_face_id_logs(request):
    """Barcha xodimlarning Face ID davomat jurnali (Faqat Admin)."""
    filter_period = request.GET.get('period', 'today')
    filter_date_str = request.GET.get('date', '')
    filter_role = request.GET.get('role', 'all')
    filter_status = request.GET.get('status', 'all')

    today = date.today()
    logs_qs = FaceIDLog.objects.select_related('user').all()

    if filter_date_str:
        try:
            filter_date = datetime.strptime(filter_date_str, '%Y-%m-%d').date()
            logs_qs = logs_qs.filter(date=filter_date)
            filter_period = 'custom'
        except Exception:
            filter_date_str = today.strftime('%Y-%m-%d')
            logs_qs = logs_qs.filter(date=today)
            filter_period = 'today'
    elif filter_period == 'today':
        logs_qs = logs_qs.filter(date=today)
        filter_date_str = today.strftime('%Y-%m-%d')
    elif filter_period == 'month':
        logs_qs = logs_qs.filter(date__year=today.year, date__month=today.month)

    if filter_role in ('manager', 'teacher'):
        logs_qs = logs_qs.filter(user__role=filter_role)
    if filter_status == 'on_time':
        logs_qs = logs_qs.filter(is_late=False)
    elif filter_status == 'late':
        logs_qs = logs_qs.filter(is_late=True)

    logs_list = list(logs_qs.order_by('-date', '-check_in_time'))
    total_scans = len(logs_list)
    late_scans = sum(1 for l in logs_list if l.is_late)
    total_fines = sum(l.fine_amount_deducted for l in logs_list)

    try:
        import core.face_service as fs
        has_cv = getattr(fs, 'HAS_OPENCV', False)
    except Exception:
        fs, has_cv = None, False

    staff_status_list = (
        _serialize_staff_face_status(Manager.objects.select_related('user').all(), 'manager', 'Menejer', has_cv, fs)
        + _serialize_staff_face_status(Teacher.objects.select_related('user').all(), 'teacher', "O'qituvchi", has_cv, fs)
    )
    ready_staff_count = sum(1 for s in staff_status_list if s['face_detected'])

    return render(request, 'admin_face_id_logs.html', {
        'logs': logs_list,
        'filter_period': filter_period,
        'filter_date': filter_date_str,
        'filter_role': filter_role,
        'filter_status': filter_status,
        'total_scans': total_scans,
        'late_scans': late_scans,
        'ontime_scans': total_scans - late_scans,
        'total_fines': total_fines,
        'staff_status_list': staff_status_list,
        'ready_staff_count': ready_staff_count,
        'missing_staff_count': len(staff_status_list) - ready_staff_count,
    })


@admin_only_required
def admin_face_id_sync_staff(request):
    """Barcha mavjud xodimlarni Face ID ga bir marta bosish bilan sinxronlash (get_or_create va try-except bilan)."""
    settings_obj = FaceIDSettings.get_settings()
    synced_count = created_profiles = 0

    for role, model_cls in (
        ('manager', Manager),
        ('teacher', Teacher),
    ):
        for u in CustomUser.objects.filter(role=role):
            try:
                staff_obj, created = model_cls.objects.get_or_create(
                    user=u,
                    defaults={
                        'first_name': u.first_name or u.username,
                        'last_name': u.last_name or '',
                        'phone': '',
                        'generated_login': u.username,
                        'generated_password': '***',
                        'work_start_time': settings_obj.default_work_start_time,
                        'work_end_time': settings_obj.default_work_end_time,
                        'late_fine_rate': settings_obj.default_late_fine_rate,
                    },
                )
                if created:
                    created_profiles += 1
            except Exception:
                continue

        for staff in model_cls.objects.all():
            try:
                updated_fields = []
                if not staff.work_start_time:
                    staff.work_start_time = settings_obj.default_work_start_time
                    updated_fields.append('work_start_time')
                if not staff.work_end_time:
                    staff.work_end_time = settings_obj.default_work_end_time
                    updated_fields.append('work_end_time')
                if updated_fields:
                    staff.save(update_fields=updated_fields)
                synced_count += 1
            except Exception:
                continue

    emb_stats = sync_all_face_embeddings(force=True)
    linked_faces = emb_stats.get('teachers', 0) + emb_stats.get('managers', 0)
    errors_count = emb_stats.get('errors', 0)

    messages.success(
        request,
        f"Face ID sinxronizatsiyasi muvaffaqiyatli: {synced_count} ta xodim sozlamalari yangilandi, "
        f"{linked_faces} ta yuz biometrik shabloni serverda qayta hisoblandi, "
        f"{created_profiles} ta yetishmayotgan profil yaratildi"
        + (f" ({errors_count} ta rasmda xatolik o'tkazib yuborildi)." if errors_count else "."),
    )
    return redirect('admin_face_id_logs')


@admin_only_required
def admin_face_id_upload_photo(request, role, id):
    """Admin orqali xodimga Face ID yuz fotosurati va face-api.js descriptorini biriktirish."""
    model_map = {'teacher': Teacher, 'manager': Manager}
    if role not in model_map:
        return JsonResponse({'success': False, 'message': "Noto'g'ri lavozim"}, status=400)
    staff = get_object_or_404(model_map[role], id=id)

    descriptor_raw = request.POST.get('descriptor')
    photo_file = request.FILES.get('photo')
    image_base64 = request.POST.get('image_base64')
    if not photo_file and image_base64:
        photo_file, _, b64_err = decode_base64_image(image_base64)
        if not photo_file:
            return JsonResponse({'success': False, 'message': f"Suratni o'qishda xatolik: {b64_err}"}, status=400)

    if photo_file:
        opt_photo, err = optimize_student_photo(photo_file)
        if not opt_photo:
            return JsonResponse({'success': False, 'message': f"Rasmni qayta ishlashda xatolik: {err}"}, status=400)
        staff.photo = opt_photo
        staff.save()

    face_ok = False
    if descriptor_raw:
        face_ok = save_staff_descriptor(staff, descriptor_raw)
    else:
        face_ok = sync_staff_face_embedding(staff)

    msg = f"{staff.first_name} {staff.last_name} uchun Face ID yuz fotosurati muvaffaqiyatli biriktirildi!"
    msg += (
        " (128-o'lchamli yuz deskriptori bazaga saqlandi ✅)."
        if face_ok else
        " (Rasm saqlandi — terminal ochilganda face-api.js avtomatik deskriptor yaratadi)."
    )

    if _is_ajax(request):
        return JsonResponse({
            'success': True,
            'message': msg,
            'photo_url': staff.photo.url if staff.photo else '',
            'face_detected': face_ok,
        })

    messages.success(request, msg)
    return redirect('admin_face_id_logs')


@admin_only_required
def admin_face_id_settings(request):
    """Faqat Admin uchun Face ID va Geolokatsiya sozlamalari paneli."""
    settings_obj = FaceIDSettings.get_settings()
    if request.method == 'POST':
        lat_val = request.POST.get('latitude', '').strip()
        lon_val = request.POST.get('longitude', '').strip()
        radius_val = request.POST.get('allowed_radius_meters', '').strip()
        fine_rate_val = request.POST.get('default_late_fine_rate', '').strip()
        try:
            if lat_val:
                settings_obj.latitude = float(lat_val.replace(',', '.'))
            if lon_val:
                settings_obj.longitude = float(lon_val.replace(',', '.'))
            if radius_val:
                settings_obj.allowed_radius_meters = int(float(radius_val.replace(',', '.')))
            if fine_rate_val:
                settings_obj.default_late_fine_rate = _safe_decimal(fine_rate_val, '2.0')
            settings_obj.default_work_start_time = request.POST.get('default_work_start_time', '09:00').strip() or '09:00'
            settings_obj.default_work_end_time = request.POST.get('default_work_end_time', '18:00').strip() or '18:00'
            settings_obj.is_geofencing_enabled = request.POST.get('is_geofencing_enabled') == 'on'
            settings_obj.save()
            messages.success(request, "Face ID va Geolokatsiya sozlamalari muvaffaqiyatli yangilandi! ✅")
            return redirect('admin_face_id_settings')
        except Exception as e:
            messages.error(request, f"Saqlashda xatolik yuz berdi: {e}")

    return render(request, 'admin_face_id_settings.html', {'settings': settings_obj})


# ─────────────────────────────────────────────────────────────
# 4. FRONTEND (face-api.js) TERMINAL VA DESCRIPTOR API BACKEND
# ─────────────────────────────────────────────────────────────

def _get_staff_photos_for_face_api() -> str:
    """
    Frontend (face-api.js) bazada rasmi bor, ammo hali 128-D face-api.js deskriptori
    saqlanmagan xodimlarning rasmidan avtomatik descriptor yaratib olishi uchun JSON ro'yxat.
    """
    items = []
    for role, qs in (
        ('teacher', Teacher.objects.filter(is_active=True)),
        ('manager', Manager.objects.filter(is_active=True)),
    ):
        for s in qs:
            if s.photo:
                try:
                    items.append({
                        'role': role,
                        'id': s.id,
                        'name': f"{s.first_name} {s.last_name}".strip(),
                        'photo_url': s.photo.url,
                        'has_descriptor': parse_descriptor(s.face_embedding) is not None,
                    })
                except Exception:
                    continue
    return json.dumps(items)


@csrf_exempt
def face_id_terminal(request):
    """Markaz binosida xodimlar yuzini face-api.js orqali skanerlovchi terminal sahifasi."""
    return render(request, 'face_id_terminal.html', {
        'settings': FaceIDSettings.get_settings(),
        'teachers': Teacher.objects.filter(is_active=True).select_related('user'),
        'managers': Manager.objects.filter(is_active=True).select_related('user'),
        'staff_photos_json': _get_staff_photos_for_face_api(),
    })


@csrf_exempt
def face_id_check_in(request):
    """Mobil va veb qurilmalardan face-api.js orqali davomat qilish sahifasi."""
    return render(request, 'face_id_mobile_checkin.html', {
        'settings': FaceIDSettings.get_settings(),
        'staff_photos_json': _get_staff_photos_for_face_api(),
    })


def _resolve_target_staff_user(identifier):
    """user_id yoki employee_id orqali faol xodimning CustomUser obyektini topadi."""
    if not identifier:
        return None
    try:
        u = CustomUser.objects.select_related('teacher_profile', 'manager_profile').filter(
            id=identifier,
            is_active=True,
            role__in=('teacher', 'manager', 'admin'),
        ).first()
        if u and get_user_staff_profile(u):
            return u
        t_obj = Teacher.objects.select_related('user').filter(id=identifier, is_active=True).first()
        if t_obj and t_obj.user:
            return t_obj.user
        m_obj = Manager.objects.select_related('user').filter(id=identifier, is_active=True).first()
        if m_obj and m_obj.user:
            return m_obj.user
    except Exception:
        pass
    return None


@csrf_exempt
def face_id_verify_api(request):
    """
    Frontend'dagi face-api.js dan kelgan 128 o'lchamli `descriptor` (massiv) ni qabul qilib,
    bazadagi xodimlarning descriptor'lari bilan Evklid masofasi (Euclidean distance) orqali
    solishtiradigan va davomatni avtomatik yozadigan yengil Django API View.

    JSON javob formati:
      - Mos kelsa: {"success": true, "name": "Xodim ismi", "employee_name": "Xodim ismi", "message": "Xodim keldi deb belgilandi"}
      - Topilmasa: {"success": false, "message": "Yuz topilmadi yoki xodim ro'yxatda yo'q"}
    """
    if request.method != 'POST':
        return JsonResponse({
            'success': False,
            'message': "Faqat POST so'rovi qabul qilinadi",
        }, status=405)

    payload = {}
    if request.body:
        try:
            payload = json.loads(request.body.decode('utf-8'))
        except Exception:
            payload = {}

    # Maxsus amal: Frontend (face-api.js) xodimning profil rasmidan 128-D descriptor ajratib bazaga saqlaganda
    action = payload.get('action') or request.POST.get('action')
    if action == 'save_descriptor':
        role = payload.get('role') or request.POST.get('role')
        staff_id = payload.get('id') or request.POST.get('id')
        desc = payload.get('descriptor') or request.POST.get('descriptor')
        model_map = {'teacher': Teacher, 'manager': Manager}
        if role in model_map and staff_id and desc:
            staff_obj = model_map[role].objects.filter(id=staff_id).first()
            if staff_obj and save_staff_descriptor(staff_obj, desc):
                return JsonResponse({'success': True, 'message': "Deskriptor saqlandi"})
        return JsonResponse({'success': False, 'message': "Deskriptorni saqlab bo'lmadi"}, status=400)

    descriptor_data = (
        payload.get('descriptor')
        or payload.get('embedding')
        or request.POST.get('descriptor')
        or request.POST.get('embedding')
    )
    img_data = (
        payload.get('image')
        or payload.get('photo')
        or request.POST.get('image')
        or request.POST.get('photo')
    )
    file_obj = request.FILES.get('photo') or request.FILES.get('image')
    employee_id = (
        payload.get('employee_id')
        or payload.get('user_id')
        or request.POST.get('employee_id')
        or request.POST.get('user_id')
    )

    lat = payload.get('latitude') if 'latitude' in payload else request.POST.get('latitude')
    lon = payload.get('longitude') if 'longitude' in payload else request.POST.get('longitude')
    is_kiosk = str(payload.get('is_kiosk', request.POST.get('is_kiosk', '0'))).lower() in ('1', 'true')
    bypass_geo = (
        str(payload.get('bypass_geo', request.POST.get('bypass_geo', '0'))).lower() in ('1', 'true')
        or (is_kiosk and not lat and not lon)
    )

    # 1. Frontend yuborgan 128-D descriptor massivini tekshiramiz
    scan_vec = parse_descriptor(descriptor_data)
    if scan_vec is None:
        return JsonResponse({
            'success': False,
            'detected': False,
            'error_type': 'no_face',
            'message': "Yuz topilmadi yoki xodim ro'yxatda yo'q",
        }, status=200)

    # Ixtiyoriy: jurnalga rasm biriktirish uchun snapshot yuborilgan bo'lsa
    scan_file = file_obj
    if not scan_file and img_data:
        scan_file, _, _ = decode_base64_image(img_data)

    target_user = None
    similarity_score = 95.0

    # 2. Bazadagi barcha xodimlarning 128-D descriptorlari bilan Euclidean distance bo'yicha solishtirish
    match = match_staff_by_face(scan_descriptor=scan_vec)
    if match:
        _, matched_staff, similarity_score, _ = match
        target_user = matched_staff.user
    elif employee_id:
        candidate_user = _resolve_target_staff_user(employee_id)
        if candidate_user:
            staff_prof = get_user_staff_profile(candidate_user)
            staff_vec = get_or_create_face_embedding(staff_prof) if staff_prof else None
            if staff_vec is not None:
                score, is_match, _ = compare_embeddings(staff_vec, scan_vec)
                if is_match:
                    target_user = candidate_user
                    similarity_score = score

    # 3. Agar mos xodim topilmasa
    if not target_user:
        return JsonResponse({
            'success': False,
            'detected': True,
            'error_type': 'unknown_face',
            'message': "Yuz topilmadi yoki xodim ro'yxatda yo'q",
        }, status=200)

    # 4. Davomatni (FaceIDLog) avtomatik yozish va JSON javob qaytarish
    res = process_face_check_in(
        user=target_user,
        scan_file=scan_file,
        lat=lat,
        lon=lon,
        bypass_geo=bypass_geo,
        pre_verified=True,
        match_score=similarity_score,
    )

    if not res.get('success'):
        return JsonResponse({
            'success': False,
            'is_geo_error': res.get('is_geo_error', False),
            'message': res.get('message') or "Yuz topilmadi yoki xodim ro'yxatda yo'q",
        }, status=200)

    staff_obj, staff_photo, _ = get_user_face_data(target_user)
    if staff_obj and staff_photo:
        try:
            res['photo_url'] = staff_photo.url
        except Exception:
            pass

    if not is_kiosk and getattr(target_user, 'role', '') in ('teacher', 'manager'):
        if not getattr(request, 'user', None) or not request.user.is_authenticated:
            try:
                login(request, target_user)
            except Exception:
                pass
        if target_user.role == 'teacher':
            t_prof = get_user_staff_profile(target_user)
            res['redirect_url'] = '/teacher/dashboard/' if (t_prof and t_prof.groups.exists()) else '/staff/my-logs/'
        elif target_user.role == 'manager':
            res['redirect_url'] = '/manager/dashboard/'

    return JsonResponse(res)


# ─────────────────────────────────────────────────────────────
# 5. OYLIK MAOSH, BONUSLAR, AUDIT VA HISOBOTLAR
# ─────────────────────────────────────────────────────────────

@admin_only_required
def admin_salary_report(request):
    """Barcha xodimlarning oyligi, kechikish jarimalari, bonus va avanslarini hisoblash sahifasi."""
    month_str, year, month = normalize_month_str(request.GET.get('month'))
    all_managers = Manager.objects.filter(is_active=True).select_related('user')
    all_teachers = Teacher.objects.filter(is_active=True).select_related('user')

    reports = []
    total_base = total_deductions = total_advances = total_bonuses = total_net = Decimal(0)
    total_lates = 0

    for staff, role_title in (
        *[(m, "Menejer") for m in all_managers],
        *[(t, "O'qituvchi / Tarbiyachi") for t in all_teachers],
    ):
        summary = calculate_staff_salary_summary(staff, year, month)
        summary['role_title'] = role_title
        reports.append(summary)
        total_base += summary['base_salary']
        total_deductions += summary['total_fine_amount']
        total_advances += summary['total_advance_amount']
        total_bonuses += summary['total_bonus_amount']
        total_net += summary['net_salary']
        total_lates += summary['total_late_count']

    return render(request, 'admin_salary_report.html', {
        'reports': reports,
        'current_month': month_str,
        'year': year,
        'month': month,
        'total_base_payroll': total_base,
        'total_deductions': total_deductions,
        'total_advances_center': total_advances,
        'total_bonuses_center': total_bonuses,
        'total_net_payroll': total_net,
        'total_lates_center': total_lates,
    })


@admin_only_required
def admin_bonus_save(request):
    """Direktor tomonidan xodimga bonus qo'shish (Faqat Admin)."""
    if request.method == 'POST':
        user_id = request.POST.get('user_id')
        month, _, _ = normalize_month_str(request.POST.get('month'))
        amount_raw = request.POST.get('amount', '0').strip().replace(',', '').replace(' ', '')
        reason = request.POST.get('reason', "Erta kelganlik / Faollik uchun bonus").strip()
        try:
            target_user = CustomUser.objects.get(id=user_id)
            amount = Decimal(amount_raw)
            if amount > 0:
                actor = request.user if getattr(request, 'user', None) and request.user.is_authenticated else None
                StaffBonus.objects.create(
                    user=target_user, month=month, date=date.today(), amount=amount, reason=reason, created_by=actor
                )
                AuditLog.objects.create(
                    action_type='bonus',
                    actor=actor,
                    actor_name="Direktor",
                    title=f"Bonus berildi: {target_user.staff_full_name} (+{amount:,.0f} so'm)",
                    details=f"Sabab: {reason}. Oy: {month}.",
                )
                messages.success(request, f"{target_user.staff_full_name} ga {amount:,.0f} so'm bonus muvaffaqiyatli qo'shildi! 🎉")
            else:
                messages.error(request, "Bonus summasi 0 dan katta bo'lishi kerak!")
        except Exception as e:
            messages.error(request, f"Bonus qo'shishda xatolik: {e}")

    return redirect(f"{request.META.get('HTTP_REFERER', '/admin/salary/report/')}?month={request.POST.get('month', '')}")


@admin_only_required
def admin_bonus_delete(request, id):
    """Berilgan bonusni bekor qilish (Faqat Admin)."""
    bonus = get_object_or_404(StaffBonus, id=id)
    u_name, amt, m = bonus.user.staff_full_name, bonus.amount, bonus.month
    bonus.delete()
    AuditLog.objects.create(
        action_type='bonus',
        actor=request.user if getattr(request, 'user', None) and request.user.is_authenticated else None,
        actor_name="Direktor",
        title=f"Bonus bekor qilindi: {u_name} (-{amt:,.0f} so'm)",
        details=f"Oy: {m}.",
    )
    messages.info(request, f"{u_name} ning {amt:,.0f} so'm bonusi bekor qilindi.")
    return redirect(request.META.get('HTTP_REFERER', '/admin/salary/report/'))


@admin_only_required
def admin_audit_logs(request):
    """Tizim Tarixi (Audit Log): Faqat Admin uchun."""
    filter_action = request.GET.get('action', 'all')
    filter_date = request.GET.get('date', '')
    search_q = request.GET.get('q', '').strip()

    logs_qs = AuditLog.objects.select_related('actor').all()
    if filter_action and filter_action != 'all':
        logs_qs = logs_qs.filter(action_type=filter_action)
    if filter_date:
        try:
            logs_qs = logs_qs.filter(created_at__date=datetime.strptime(filter_date, '%Y-%m-%d').date())
        except Exception:
            pass
    if search_q:
        logs_qs = logs_qs.filter(
            Q(title__icontains=search_q) | Q(details__icontains=search_q) | Q(actor_name__icontains=search_q)
        )

    return render(request, 'admin_audit_logs.html', {
        'logs': logs_qs[:150],
        'filter_action': filter_action,
        'filter_date': filter_date,
        'search_q': search_q,
    })


@admin_only_required
def admin_salary_export_excel(request):
    """Oylik va jarimalar hisobini Excel (.xlsx) faylga eksport qilish (Faqat Admin)."""
    month_str, year, month = normalize_month_str(request.GET.get('month'))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Oylik_{month_str}"

    headers = [
        "№", "F.I.Sh.", "Lavozim", "Telefon", "Asosiy oylik (so'm)",
        "Ish kunlari", "Kechikishlar soni", "Kechikilgan daqiqa",
        "Jarima foizi (%)", "Jarima summasi (so'm)", "Avans (so'm)", "Bonus (so'm)", "Qo'lga tegadigan oylik (so'm)",
    ]
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = openpyxl.styles.Font(bold=True, color="FFFFFF", size=11)
        cell.fill = openpyxl.styles.PatternFill(start_color="1E1B4B", end_color="1E1B4B", fill_type="solid")
        cell.alignment = openpyxl.styles.Alignment(horizontal="center", vertical="center")

    staff_list = list(Manager.objects.filter(is_active=True)) + list(Teacher.objects.filter(is_active=True))
    for idx, s in enumerate(staff_list, 1):
        summary = calculate_staff_salary_summary(s, year, month)
        ws.append([
            idx,
            f"{s.first_name} {s.last_name}",
            s.role_title,
            s.phone or "-",
            float(summary['base_salary']),
            summary['total_days'],
            summary['total_late_count'],
            summary['total_late_minutes'],
            float(summary['total_fine_percent']),
            float(summary['total_fine_amount']),
            float(summary['total_advance_amount']),
            float(summary['total_bonus_amount']),
            float(summary['net_salary']),
        ])

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="HUMO_Oylik_va_Jarimalar_{month_str}.xlsx"'
    wb.save(response)
    return response


# ─────────────────────────────────────────────────────────────
# 6. MENEJER PANELI, SHAXSIY JURNAL VA AVANS SO'ROVLARI
# ─────────────────────────────────────────────────────────────

@login_required(login_url='/manager/login/')
def manager_dashboard(request):
    """Menejer Paneli: Shaxsiy profil, Face ID jurnali, oylik maosh va markaz guruhlari."""
    if request.user.role != 'manager':
        return redirect('home')

    manager, _ = Manager.objects.get_or_create(
        user=request.user,
        defaults={
            'first_name': request.user.first_name or request.user.username,
            'last_name': request.user.last_name or '',
            'generated_login': request.user.username,
            'generated_password': '***',
        },
    )
    today = date.today()
    return render(request, 'manager_dashboard.html', {
        'manager': manager,
        'today_log': FaceIDLog.objects.filter(user=request.user, date=today).first(),
        'salary_summary': calculate_staff_salary_summary(manager, today.year, today.month),
        'face_logs': FaceIDLog.objects.filter(user=request.user).order_by('-date')[:30],
        'groups': Group.objects.select_related('teacher').prefetch_related('time_slots__students'),
    })


@login_required
def staff_my_face_logs(request):
    """Har bir xodim uchun shaxsiy Face ID kelib-ketish tarixi va avans so'rovlari."""
    staff = request.user.staff_profile
    if not staff:
        messages.error(request, "Xodim profili biriktirilmagan.")
        return redirect('home')

    month_str, year, month = normalize_month_str(request.GET.get('month'))
    return render(request, 'staff_my_face_logs.html', {
        'staff': staff,
        'salary_summary': calculate_staff_salary_summary(staff, year, month),
        'current_month': month_str,
        'my_advances': AdvanceRequest.objects.filter(user=request.user).order_by('-created_at')[:20],
    })


@admin_required
def admin_advance_list(request):
    """Xodimlar tomonidan yuborilgan avans so'rovlari ro'yxati (Admin va Menejer uchun)."""
    filter_status = request.GET.get('status', 'all')
    filter_month = request.GET.get('month', '')

    advances_qs = AdvanceRequest.objects.select_related('user', 'reviewed_by').all()
    if filter_status and filter_status != 'all':
        advances_qs = advances_qs.filter(status=filter_status)
    if filter_month:
        advances_qs = advances_qs.filter(month=filter_month)

    agg = AdvanceRequest.objects.aggregate(
        total=Count('id'),
        pending=Count('id', filter=Q(status='pending')),
        approved=Count('id', filter=Q(status='approved')),
        rejected=Count('id', filter=Q(status='rejected')),
        approved_sum=Sum('amount', filter=Q(status='approved')),
    )
    stats = {
        'total': agg['total'] or 0,
        'pending': agg['pending'] or 0,
        'approved': agg['approved'] or 0,
        'rejected': agg['rejected'] or 0,
        'approved_sum': agg['approved_sum'] or Decimal(0),
    }

    return render(request, 'admin_advances.html', {
        'advances': advances_qs,
        'filter_status': filter_status,
        'filter_month': filter_month,
        'stats': stats,
    })


@admin_required
def admin_advance_action(request, id):
    """Avans so'rovini tasdiqlash yoki rad etish (POST)."""
    if request.method != 'POST':
        return redirect('admin_advance_list')

    advance = get_object_or_404(AdvanceRequest, id=id)
    action = request.POST.get('action')
    comment = request.POST.get('admin_comment', '').strip()

    action_map = {
        'approve': ('approved', "tasdiqlandi ✅"),
        'reject': ('rejected', "rad etildi ❌"),
    }
    if action not in action_map:
        messages.error(request, "Noto'g'ri buyruq yuborildi.")
        return redirect('admin_advance_list')

    advance.status, status_text = action_map[action]
    reviewer_user = request.user if getattr(request, 'user', None) and request.user.is_authenticated else None
    advance.admin_comment = comment
    advance.reviewed_by = reviewer_user
    advance.reviewed_at = timezone.now()
    advance.save()

    AuditLog.objects.create(
        action_type='advance',
        actor=reviewer_user,
        actor_name="Admin" if is_admin_user(request) else "Menejer",
        title=f"Avans {status_text}: {advance.user.staff_full_name} ({advance.amount:,.0f} so'm)",
        details=f"Oy: {advance.month}. Izoh: {comment or 'Izoh kiritilmagan'}.",
    )
    messages.success(request, f"{advance.user.staff_full_name} ning {advance.amount:,.0f} so'm miqdoridagi avansi {status_text}!")
    return redirect('admin_advance_list')


@login_required
def staff_advance_request(request):
    """Xodim (O'qituvchi yoki Menejer) shaxsiy kabinetidan avans so'rash."""
    if request.method == 'POST':
        amount_raw = request.POST.get('amount', '0').strip().replace(',', '').replace(' ', '')
        reason = request.POST.get('reason', '').strip()
        month, _, _ = normalize_month_str(request.POST.get('month'))
        try:
            amount = Decimal(amount_raw)
            if amount <= 0:
                messages.error(request, "Avans summasi 0 dan katta bo'lishi kerak!")
                return redirect(request.META.get('HTTP_REFERER', '/staff/my-logs/'))

            AdvanceRequest.objects.create(
                user=request.user, amount=amount, reason=reason, month=month, status='pending'
            )
            AuditLog.objects.create(
                action_type='advance',
                actor=request.user,
                actor_name=request.user.staff_full_name,
                title=f"Avans so'rovi yuborildi: {request.user.staff_full_name} ({amount:,.0f} so'm)",
                details=f"Oy: {month}. Sabab: {reason or 'Izoh kiritilmagan'}.",
            )
            messages.success(
                request,
                f"{amount:,.0f} so'm miqdoridagi avans so'rovi adminga muvaffaqiyatli yuborildi! "
                "Tasdiqlangach oylikdan avtomatik chegiriladi.",
            )
        except Exception as e:
            messages.error(request, f"Avans so'rovida xatolik: {e}")

    return redirect(request.META.get('HTTP_REFERER', '/staff/my-logs/'))

