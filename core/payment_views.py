"""
HUMO O'quv markazi — Payment Views & Services (payment_views.py)
To'lovlar, qarzdorlar filtri, ORM yordamchi funksiyalar va Excel eksport.
"""
import calendar
import re
from datetime import date, datetime

try:
    from dateutil.relativedelta import relativedelta
except ImportError:
    relativedelta = None

import openpyxl
from django.contrib import messages
from django.db.models import Exists, OuterRef, Q
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render

from .auth_views import _admin_only, admin_required
from .models import AuditLog, DeletedStudent, Group, MonthlyPayment, Student, Teacher, TimeSlot
from .payment_excel import generate_payments_excel
from .teacher_views import _populate_excel_sheet


# ─────────────────────────────────────────────────────────────
# 1. UMUMIY YORDAMCHI FUNKSIYALAR (UTILS & ORM QUERYSET HELPERS)
# ─────────────────────────────────────────────────────────────

def _is_ajax(request) -> bool:
    """So'rov AJAX / Fetch orqali kelganligini tekshiradi."""
    return (
        request.headers.get('x-requested-with') == 'XMLHttpRequest'
        or request.POST.get('ajax') == '1'
        or request.POST.get('is_ajax') in ('1', 'true', 'True')
        or 'application/json' in request.headers.get('Accept', '')
    )


def normalize_month_str(req_month=None) -> tuple[str, int, int]:
    """
    Har qanday oy satrini ('2026-9', ' 2026-10 ', None, '') qat'iy 'YYYY-MM'
    formatiga keltiradi va (month_str, year, month) qaytaradi.
    """
    today = date.today()
    try:
        y, m = map(int, str(req_month or '').strip().split('-')[:2])
        if not (1 <= m <= 12):
            raise ValueError
    except Exception:
        y, m = today.year, today.month
    return f"{y:04d}-{m:02d}", y, m


def _clean_amount(val) -> int:
    """
    To'lov qiymatidan sof raqamni ajratib oladi.
    500000, '500 000', '500,000', '500000.00', '500 000 so'm' -> int.
    Matn ('mahalla', 'B.U') yoki bo'sh bo'lsa -> 0.
    """
    if not val:
        return 0
    if isinstance(val, (int, float)):
        return int(val)
    s = str(val).strip().replace('\xa0', ' ')
    if not s or s.lower() in ('0', 'none', '-', 'null', '0.0', '0.00'):
        return 0
    for w in ("so'm", "som", "so`m", "uzs"):
        s = re.sub(re.escape(w), '', s, flags=re.IGNORECASE).strip()
    if '.' in s and s.replace('.', '').replace(' ', '').replace(',', '').isdigit():
        parts = s.replace(' ', '').replace(',', '').split('.')
        # Agar '500.000' (minglik ajratuvchi) bo'lsa -> 500000, '500000.00' (tiyin) bo'lsa -> 500000
        if len(parts) == 2 and len(parts[1]) <= 2:
            try:
                return int(float(s.replace(' ', '').replace(',', '')))
            except Exception:
                pass
    digits = "".join(c for c in s if c.isdigit())
    return int(digits) if digits else 0


def _is_text_payment(val) -> bool:
    """
    To'lov katagiga matn kiritilganligini tekshiradi ('mahalla', 'B.U', 'tekin', 'bepul').
    Raqamli formatlar ('500 000', '500,000', '500000.00', '0', '0.00') matn EMAS.
    """
    if not val:
        return False
    s = str(val).strip().replace('\xa0', ' ')
    if not s or s.lower() in ('0', '-', 'none', 'null', '0.0', '0.00'):
        return False
    for w in ("so'm", "som", "so`m", "uzs"):
        s = re.sub(re.escape(w), '', s, flags=re.IGNORECASE).strip()
    digits = s.replace(' ', '').replace("'", '').replace(',', '').replace('.', '')
    return bool(digits) and not digits.isdigit()


def _is_mahalla_student(student, pay=None) -> bool:
    """
    O'quvchi mahalla toifasiga kirishini tekshiradi:
      1. Student.is_local == True yoki Student.is_mahalla == True
      2. To'lov summasiga (amount_paid) yoki izohiga (notes) 'mahalla' yozilgan
    """
    if getattr(student, 'is_local', False) or getattr(student, 'is_mahalla', False):
        return True
    if pay:
        raw_amt = str(getattr(pay, 'amount_paid', '') or '').strip().lower()
        raw_notes = str(getattr(pay, 'notes', '') or '').strip().lower()
        if 'mahalla' in raw_amt or 'mahalla' in raw_notes:
            return True
    return False


def _is_free_student(student, pay=None) -> bool:
    """O'quvchi bepul / imtiyozli ('bepul', 'tekin', 'free', 'grant') ekanligini tekshiradi."""
    if getattr(student, 'is_free', False):
        return True
    if pay:
        raw_amt = str(getattr(pay, 'amount_paid', '') or '').strip().lower()
        raw_notes = str(getattr(pay, 'notes', '') or '').strip().lower()
        free_keywords = ('bepul', 'tekin', 'free', 'grant')
        if any(k in raw_amt or k in raw_notes for k in free_keywords):
            return True
    return False


def _payment_paid_info(pay) -> tuple[str, int, bool, bool]:
    """
    To'lov yozuvi bo'yicha YAGONA va qat'iy "to'langanmi?" tekshiruvi.
    Qaytaradi: (raw_paid, num_paid, has_text_payment, is_paid)
      is_paid = True, agar:
        - amount_paid > 0 (raqamli summa), YOKI
        - amount_paid ga matn yozilgan ('B.U', 'tekin' va h.k.), YOKI
        - status == 'paid'
    """
    if not pay:
        return '', 0, False, False
    raw_paid = str(pay.amount_paid or '').strip()
    num_paid = _clean_amount(raw_paid)
    has_text = _is_text_payment(raw_paid)
    is_paid = (num_paid > 0) or has_text or (pay.status == 'paid')
    return raw_paid, num_paid, has_text, is_paid


def _normalize_phone(phone_str) -> str:
    if not phone_str:
        return ""
    digits = re.sub(r'\D', '', str(phone_str))
    return digits[-9:] if len(digits) >= 9 else digits


def _get_student_person_key(student) -> str:
    """Jismoniy bitta o'quvchini telefon va ism orqali unikal aniqlash kaliti."""
    clean_p = _normalize_phone(getattr(student, 'phone_1', ''))
    fn = re.sub(r'[^a-zA-Z0-9]', '', str(getattr(student, 'first_name', '') or '').lower())
    if clean_p and fn:
        return f"p_{clean_p}_{fn}"
    if clean_p:
        return f"p_{clean_p}"
    ln = re.sub(r'[^a-zA-Z0-9]', '', str(getattr(student, 'last_name', '') or '').lower())
    return f"name_{fn}_{ln}" if (fn or ln) else f"id_{student.id}"


def heal_student_group_links():
    """
    Ma'lumotlar bazasidagi o'quvchi-guruh bog'lanishlaridagi nomuvofiqliklarni avtomatik tuzatadi:
      1. Guruhga (`time_slot`) biriktirilgan, ammo arxivda (`DeletedStudent`) bo'lmagan holda
         `is_active=False` yoki `left_date` qolib ketgan o'quvchilarni `is_active=True, left_date=None` qiladi.
      2. Faol (`is_active=True`) o'quvchilarning qolib ketgan `left_date` va `DeletedStudent` qoldiqlarini tozalaydi.
      3. Faol (`is_active=True`), ammo `time_slot=None` bo'lib qolgan o'quvchilarni (agar `DeletedStudent`
         yoki guruh tarixida `time_slot_id` / `group_name` mavjud bo'lsa) o'z guruhining `TimeSlot`iga qayta ulaydi.
    """
    deleted_student_ids = set(
        DeletedStudent.objects.filter(student_id__isnull=False).values_list('student_id', flat=True)
    )

    # 1. Guruhga biriktirilgan va arxivlanmagan o'quvchi nofaol qolib ketgan bo'lsa -> faollashtiramiz
    Student.objects.filter(
        time_slot__isnull=False,
        time_slot__group__isnull=False,
        is_active=False,
    ).exclude(id__in=deleted_student_ids).update(is_active=True, left_date=None)

    # 2. Faol o'quvchilarda left_date yoki DeletedStudent qoldig'i bo'lsa -> tozalaymiz
    active_ids = list(Student.objects.filter(is_active=True).values_list('id', flat=True))
    if active_ids:
        Student.objects.filter(id__in=active_ids, left_date__isnull=False).update(left_date=None)
        if deleted_student_ids.intersection(active_ids):
            for s in Student.objects.filter(id__in=deleted_student_ids.intersection(active_ids), is_active=True, time_slot__isnull=True):
                ds = DeletedStudent.objects.filter(student_id=s.id).order_by('-deleted_at').first()
                if ds:
                    ts = None
                    if ds.time_slot_id:
                        ts = TimeSlot.objects.filter(id=ds.time_slot_id).first()
                    if not ts and ds.group_name:
                        ts = TimeSlot.objects.filter(group__name__iexact=ds.group_name.strip()).first()
                    if ts:
                        s.time_slot = ts
                        s.save(update_fields=['time_slot'])
            DeletedStudent.objects.filter(student_id__in=active_ids).delete()


def get_active_students_qs(group=None, group_id=None, search_q: str = '', require_group: bool = False):
    """
    Faol (`is_active=True`) o'quvchilar uchun markazlashgan Django ORM QuerySet.
    - Agar aniq guruh (`group` yoki `group_id`) berilsa yoki `require_group=True` bo'lsa,
      shu guruhning `TimeSlot`iga biriktirilgan faol o'quvchilarni qaytaradi.
    - Umumiy ro'yxatda (`admin_payments_all`) barcha `is_active=True` o'quvchilarni qaytaradi
      (hatto `time_slot` vaqtincha uzilib qolgan bo'lsa ham, o'quvchi yashirinib qolmaydi).
    """
    heal_student_group_links()
    qs = Student.objects.filter(is_active=True).select_related(
        'time_slot__group__teacher', 'time_slot__group'
    )
    if group is not None:
        qs = qs.filter(time_slot__group=group)
    elif group_id and str(group_id).isdigit():
        qs = qs.filter(time_slot__group_id=int(group_id))
    elif require_group:
        qs = qs.filter(time_slot__isnull=False, time_slot__group__isnull=False)
    if group is not None:
        qs = qs.filter(time_slot__group=group)
    elif group_id and str(group_id).isdigit():
        qs = qs.filter(time_slot__group_id=int(group_id))

    if search_q:
        q_filter = Q()
        for w in search_q.split():
            q_filter |= (
                Q(first_name__icontains=w)
                | Q(last_name__icontains=w)
                | Q(phone_1__icontains=w)
                | Q(phone_2__icontains=w)
            )
        qs = qs.filter(q_filter)

    return qs


def filter_students_by_payment_status(students_qs, req_month: str, status_filter: str = 'all'):
    """
    Django ORM darajasida (`Exists` va `OuterRef` orqali) o'quvchilarni
    'paid' (to'laganlar), 'unpaid' (qarzdorlar) va 'local' (mahalla) bo'yicha filtrlash.
    Shartlar:
      1. Tanlangan oyda (`req_month`) `amount_paid` bo'sh/'0' bo'lmasa YOKI `status='paid'`
         bo'lsa -> o'quvchi to'lagan (`_has_paid=True`).
      2. To'lagan o'quvchi QAT'IY ravishda `unpaid` ro'yxatidan chiqariladi (`_has_paid=False`).
      3. Faqat to'lov qilmagan, mahalla (`is_local=False`) va bepul bo'lmaganlar `unpaid` ga kiradi.
    """
    req_month, _, _ = normalize_month_str(req_month)
    empty_values = ['', '0', '0.0', '0.00', '-', 'none', 'None', 'null', 'NULL']

    paid_payment_qs = MonthlyPayment.objects.filter(
        student_id=OuterRef('pk'),
        month=req_month,
    ).filter(
        ~Q(amount_paid__in=empty_values) | Q(status='paid')
    )

    mahalla_payment_qs = MonthlyPayment.objects.filter(
        student_id=OuterRef('pk'),
        month=req_month,
    ).filter(
        Q(amount_paid__icontains='mahalla') | Q(notes__icontains='mahalla')
    )

    free_payment_qs = MonthlyPayment.objects.filter(
        student_id=OuterRef('pk'),
        month=req_month,
    ).filter(
        Q(amount_paid__icontains='bepul')
        | Q(amount_paid__icontains='tekin')
        | Q(amount_paid__icontains='free')
        | Q(notes__icontains='bepul')
        | Q(notes__icontains='tekin')
    )

    qs = students_qs.annotate(
        _has_paid=Exists(paid_payment_qs),
        _has_mahalla_pay=Exists(mahalla_payment_qs),
        _has_free_pay=Exists(free_payment_qs),
    )

    if status_filter == 'paid':
        return qs.filter(_has_paid=True)
    if status_filter == 'unpaid':
        return qs.filter(
            _has_paid=False,
            is_local=False,
            _has_mahalla_pay=False,
            _has_free_pay=False,
        )
    if status_filter == 'local':
        return qs.filter(Q(is_local=True) | Q(_has_mahalla_pay=True))
    return qs


def _fmt_money(val: int) -> str:
    return f"{val:,}".replace(',', ' ')


# ─────────────────────────────────────────────────────────────
# 2. TO'LOVLAR JADVALI VA STATISTIKA GENERATORI
# ─────────────────────────────────────────────────────────────

def _build_payments_data(students_qs, req_month, status_filter='all', is_group_specific=False):
    """
    O'quvchilar va ularning to'lovlari bo'yicha jadval satrlari hamda moliyaviy statistikani qaytaradi.
    Qat'iy qoidalar:
      1. `amount_paid > 0`, matnli to'lov yoki `status == 'paid'` bo'lsa -> FAQAT "To'lagan" ('paid').
      2. To'lagan o'quvchi qarzdorlar ('unpaid') va kechikkanlar ('overdue') ro'yxatiga MUTLAQO kirmaydi.
      3. Faqat haqiqatda to'lov qilmagan, mahalla va bepul bo'lmagan o'quvchilar qarzdor ('unpaid') hisoblanadi.
    """
    today = date.today()
    req_month, req_year, req_mon = normalize_month_str(req_month)
    status_filter = (status_filter or 'all').strip().lower()
    max_day_in_month = calendar.monthrange(req_year, req_mon)[1]
    curr_month_str = today.strftime('%Y-%m')

    heal_student_group_links()
    base_filter = {'is_active': True}
    if is_group_specific:
        base_filter['time_slot__isnull'] = False
        base_filter['time_slot__group__isnull'] = False

    students_list = list(
        students_qs.filter(**base_filter)
        .select_related('time_slot__group__teacher', 'time_slot__group')
        .distinct()
        .order_by('first_name', 'last_name', 'id')
    )
    student_ids = [s.id for s in students_list]

    payments = list(MonthlyPayment.objects.filter(student_id__in=student_ids, month=req_month))
    payment_map = {}
    for p in payments:
        prev = payment_map.get(p.student_id)
        if prev is None or (not _payment_paid_info(prev)[3] and _payment_paid_info(p)[3]):
            payment_map[p.student_id] = p

    all_rows = []
    total_rate_sum = paid_count = unpaid_count = overdue_count = local_count = free_count = debt_sum = 0

    for s in students_list:
        pay = payment_map.get(s.id)
        raw_paid, num_paid, has_text_payment, is_paid = _payment_paid_info(pay)
        is_mahalla = _is_mahalla_student(s, pay)
        is_free = not is_mahalla and _is_free_student(s, pay)

        grp = s.time_slot.group if s.time_slot else None
        rate_val = s.payment_rate or 0
        total_rate_sum += rate_val

        pay_date_obj = (
            pay.payment_date if (pay and pay.payment_date)
            else (pay.created_at.date() if (pay and pay.created_at) else None)
        )

        has_paid_val = is_paid or is_mahalla or is_free
        if pay_date_obj and has_paid_val:
            yr, mn = pay_date_obj.year, pay_date_obj.month + 1
            if mn > 12:
                mn, yr = 1, yr + 1
            due_date = date(yr, mn, min(pay_date_obj.day, calendar.monthrange(yr, mn)[1]))
        else:
            joined_d = s.joined_date or today
            due_date = date(req_year, req_mon, min(joined_d.day, max_day_in_month))

        is_overdue = False
        delay_days = 0

        if is_mahalla:
            st_code = 'local'
            debt_val = 0
            local_count += 1
            if num_paid > 0:
                paid_count += 1
                status_display = f"To'langan ({_fmt_money(num_paid)} so'm)"
            elif has_text_payment and raw_paid.lower() != 'mahalla':
                status_display = f"Mahalla ({raw_paid})"
            else:
                status_display = "Mahalla"
        elif is_free and num_paid == 0:
            st_code = 'free'
            debt_val = 0
            free_count += 1
            status_display = f"Bepul ({raw_paid})" if raw_paid else "Bepul"
        elif is_paid:
            st_code = 'paid'
            debt_val = 0
            paid_count += 1
            if num_paid > 0:
                status_display = f"To'langan ({_fmt_money(num_paid)} so'm)"
            elif has_text_payment:
                status_display = f"To'langan ({raw_paid})"
            else:
                status_display = "To'langan"
        else:
            st_code = 'unpaid'
            debt_val = rate_val
            debt_sum += debt_val
            unpaid_count += 1

            if req_month < curr_month_str:
                is_overdue = True
                delay_days = max(1, (today - due_date).days)
            elif today > due_date:
                is_overdue = True
                delay_days = (today - due_date).days

            if is_overdue:
                overdue_count += 1
                status_display = "Muddati o'tgan"
            else:
                status_display = "To'lanmagan"

        all_rows.append({
            'index': 0,
            'student': s,
            'student_id': s.id,
            'student_name': f"{s.first_name} {s.last_name}",
            'is_local': is_mahalla,
            'is_mahalla': is_mahalla,
            'is_free': is_free,
            'is_paid': is_paid,
            'group': grp,
            'group_name': grp.name if grp else "-",
            'group_id': grp.id if grp else None,
            'subject_name': grp.subject_name if grp else "-",
            'teacher_name': str(grp.teacher) if (grp and grp.teacher) else "-",
            'joined_date': s.joined_date.strftime('%d.%m.%Y') if s.joined_date else "-",
            'joined_date_iso': s.joined_date.strftime('%Y-%m-%d') if s.joined_date else "",
            'due_date': due_date,
            'due_date_str': due_date.strftime('%d.%m.%Y'),
            'is_overdue': is_overdue,
            'delay_days': delay_days,
            'delay_text': f"{delay_days} kun kechikdi" if is_overdue else "",
            'payment_date': pay_date_obj.strftime('%d.%m.%Y') if pay_date_obj else "",
            'payment_date_iso': pay_date_obj.strftime('%Y-%m-%d') if pay_date_obj else "",
            'payment_rate': rate_val,
            'payment_rate_formatted': _fmt_money(rate_val),
            'amount_paid': raw_paid,
            'amount_paid_formatted': _fmt_money(num_paid) if num_paid > 0 else (raw_paid or "-"),
            'num_paid': num_paid,
            'has_text_payment': has_text_payment,
            'status_code': st_code,
            'status_display': status_display,
            'phone': f"{s.phone_1} {s.phone_2 or ''}".strip(),
            'notes': pay.notes if pay else "",
            'payment_id': pay.id if pay else None,
            'payment_type': pay.payment_type if (pay and getattr(pay, 'payment_type', None)) else 'cash',
            'payment_type_display': (
                pay.get_payment_type_display()
                if (pay and hasattr(pay, 'get_payment_type_display'))
                else "Naqd pul (Cash)"
            ),
            'debt': debt_val,
            'debt_formatted': _fmt_money(debt_val),
        })

    # Filtr bo'yicha saralash:
    # - 'paid': faqat haqiqatda to'lov qilganlar (yoki pul to'lagan mahalla o'quvchilari)
    # - 'unpaid': faqat to'lov qilmagan, mahalla bo'lmagan va bepul bo'lmagan haqiqiy qarzdorlar
    rows = []
    for r in all_rows:
        is_truly_unpaid = (
            r['status_code'] == 'unpaid'
            and not r['is_paid']
            and not r['is_mahalla']
            and not r['is_free']
        )
        if status_filter == 'unpaid' and not is_truly_unpaid:
            continue
        if status_filter == 'overdue' and not (is_truly_unpaid and r['is_overdue']):
            continue
        if status_filter == 'paid' and not (r['status_code'] == 'paid' or (r['is_mahalla'] and r['num_paid'] > 0)):
            continue
        if status_filter == 'local' and not (r['is_mahalla'] or r['has_text_payment']):
            continue
        if status_filter == 'free' and r['status_code'] != 'free':
            continue
        rows.append(r)

    for i, r in enumerate(rows, start=1):
        r['index'] = i

    # To'lov turlari va tushumlar hisobi
    totals = {'cash': 0, 'terminal': 0, 'payme': 0, 'click': 0}
    counts = {'cash': 0, 'terminal': 0, 'payme': 0, 'click': 0}
    daily_collected = total_paid_sum = 0

    for pay in payments:
        num_paid = _clean_amount(pay.amount_paid)
        if num_paid > 0:
            total_paid_sum += num_paid
            pay_date = pay.payment_date or (pay.created_at.date() if pay.created_at else None)
            if pay_date == today:
                daily_collected += num_paid
            ptype = getattr(pay, 'payment_type', 'cash') or 'cash'
            if ptype not in totals:
                ptype = 'cash'
            totals[ptype] += num_paid
            counts[ptype] += 1

    def _pct(val):
        return round((val / total_paid_sum) * 100) if total_paid_sum > 0 else 0

    stats = {
        'total_students': len(all_rows),
        'paid_count': paid_count,
        'unpaid_count': unpaid_count,
        'overdue_count': overdue_count,
        'local_count': local_count,
        'free_count': free_count,
        'total_rate': total_rate_sum,
        'total_rate_formatted': _fmt_money(total_rate_sum),
        'total_paid': total_paid_sum,
        'total_paid_formatted': _fmt_money(total_paid_sum),
        'total_cash': totals['cash'],
        'total_cash_formatted': _fmt_money(totals['cash']),
        'total_terminal': totals['terminal'],
        'total_terminal_formatted': _fmt_money(totals['terminal']),
        'total_payme': totals['payme'],
        'total_payme_formatted': _fmt_money(totals['payme']),
        'total_click': totals['click'],
        'total_click_formatted': _fmt_money(totals['click']),
        'count_cash': counts['cash'],
        'count_terminal': counts['terminal'],
        'count_payme': counts['payme'],
        'count_click': counts['click'],
        'pct_cash': _pct(totals['cash']),
        'pct_terminal': _pct(totals['terminal']),
        'pct_payme': _pct(totals['payme']),
        'pct_click': _pct(totals['click']),
        'debt_sum': debt_sum,
        'debt_sum_formatted': _fmt_money(debt_sum),
        'daily_collected': daily_collected,
        'daily_collected_formatted': _fmt_money(daily_collected),
        'req_year': req_year,
    }
    return rows, stats


def get_payment_type_aggregates(month=None, start_date=None, end_date=None):
    """To'lov turlari (Naqd, Terminal, Payme, Click) bo'yicha tushumlar yig'indisini qaytaradi."""
    qs = MonthlyPayment.objects.all()
    if month:
        norm_m, _, _ = normalize_month_str(month)
        qs = qs.filter(month=norm_m)
    if start_date:
        qs = qs.filter(payment_date__gte=start_date)
    if end_date:
        qs = qs.filter(payment_date__lte=end_date)

    totals = {'cash': 0, 'terminal': 0, 'payme': 0, 'click': 0}
    counts = {'cash': 0, 'terminal': 0, 'payme': 0, 'click': 0}
    total_all = 0

    for p in qs:
        num = _clean_amount(p.amount_paid)
        if num > 0:
            ptype = p.payment_type if getattr(p, 'payment_type', None) in totals else 'cash'
            totals[ptype] += num
            counts[ptype] += 1
            total_all += num

    def _pct(val):
        return round((val / total_all) * 100) if total_all > 0 else 0

    return {
        'total_cash': totals['cash'],
        'total_cash_formatted': _fmt_money(totals['cash']),
        'total_terminal': totals['terminal'],
        'total_terminal_formatted': _fmt_money(totals['terminal']),
        'total_payme': totals['payme'],
        'total_payme_formatted': _fmt_money(totals['payme']),
        'total_click': totals['click'],
        'total_click_formatted': _fmt_money(totals['click']),
        'count_cash': counts['cash'],
        'count_terminal': counts['terminal'],
        'count_payme': counts['payme'],
        'count_click': counts['click'],
        'pct_cash': _pct(totals['cash']),
        'pct_terminal': _pct(totals['terminal']),
        'pct_payme': _pct(totals['payme']),
        'pct_click': _pct(totals['click']),
        'total_all': total_all,
        'total_all_formatted': _fmt_money(total_all),
    }


# ─────────────────────────────────────────────────────────────
# 3. TO'LOV SAHIFALARI VA EXCEL EKSPORT VIEW'LARI
# ─────────────────────────────────────────────────────────────

@admin_required
def admin_payments_all(request):
    """Barcha guruhlar bo'yicha umumiy to'lovlar sahifasi."""
    if not _admin_only(request):
        return redirect('home')

    req_month, _, _ = normalize_month_str(request.GET.get('month'))
    group_id = request.GET.get('group_id', '').strip()
    status_filter = (request.GET.get('status') or 'all').strip().lower()
    search_q = request.GET.get('q', '').strip()

    groups = Group.objects.select_related('teacher').order_by('name')
    students_qs = get_active_students_qs(group_id=group_id, search_q=search_q)
    rows, stats = _build_payments_data(
        students_qs, req_month, status_filter, is_group_specific=group_id.isdigit()
    )

    return render(request, 'admin_payments_all.html', {
        'rows': rows,
        'groups': groups,
        'current_month': req_month,
        'selected_group_id': int(group_id) if group_id.isdigit() else '',
        'selected_status': status_filter,
        'search_query': search_q,
        'stats': stats,
    })


@admin_required
def admin_group_payments(request, group_id):
    """Alohida bir guruh uchun to'lovlar sahifasi."""
    if not _admin_only(request):
        return redirect('home')

    group = get_object_or_404(Group.objects.select_related('teacher'), id=group_id)
    req_month, _, _ = normalize_month_str(request.GET.get('month'))
    status_filter = (request.GET.get('status') or 'all').strip().lower()
    search_q = request.GET.get('q', '').strip()

    students_qs = get_active_students_qs(group=group, search_q=search_q)
    rows, stats = _build_payments_data(students_qs, req_month, status_filter, is_group_specific=True)

    return render(request, 'admin_group_payments.html', {
        'group': group,
        'rows': rows,
        'current_month': req_month,
        'selected_status': status_filter,
        'search_query': search_q,
        'stats': stats,
    })


@admin_required
def admin_payment_save(request):
    """Admin tomonidan to'lovni saqlash / yangilash."""
    is_ajax = _is_ajax(request)
    if not _admin_only(request):
        if is_ajax:
            return JsonResponse({'status': 'error', 'msg': 'Ruxsat berilmagan'}, status=403)
        messages.error(request, "Ruxsat berilmagan!")
        return redirect('home')

    if request.method != 'POST':
        if is_ajax:
            return JsonResponse({'status': 'error', 'msg': "Noto'g'ri so'rov"}, status=400)
        return redirect('admin_payments_all')

    student_id = request.POST.get('student_id')
    if not student_id:
        if is_ajax:
            return JsonResponse({'status': 'error', 'msg': "O'quvchi ko'rsatilmadi"}, status=400)
        messages.error(request, "O'quvchi tanlanmadi!")
        return redirect('admin_payments_all')

    month, _, _ = normalize_month_str(request.POST.get('month'))
    amount_raw_str = str(request.POST.get('amount') or request.POST.get('amount_paid') or '0').strip()
    payment_type = request.POST.get('payment_type', 'cash').strip()
    valid_types = {t[0] for t in MonthlyPayment.PAYMENT_TYPE_CHOICES}
    if payment_type not in valid_types:
        payment_type = 'cash'
    notes = request.POST.get('notes', '').strip()

    student = get_object_or_404(Student, id=student_id)
    amount_num = _clean_amount(amount_raw_str)
    has_text = _is_text_payment(amount_raw_str)
    is_mahalla_user = (
        getattr(student, 'is_local', False)
        or getattr(student, 'is_mahalla', False)
        or ('mahalla' in amount_raw_str.lower())
        or ('mahalla' in notes.lower())
    )

    if amount_num > 0:
        amount_to_save = str(amount_num)
        new_status = 'paid'
    elif is_mahalla_user:
        amount_to_save = amount_raw_str if amount_raw_str else 'mahalla'
        new_status = 'paid'
    else:
        amount_to_save = amount_raw_str
        new_status = 'paid' if has_text else 'unpaid'

    joined_date_raw = request.POST.get('joined_date', '').strip()
    if joined_date_raw:
        try:
            new_joined = datetime.strptime(joined_date_raw, '%Y-%m-%d').date()
            if new_joined > date.today():
                if is_ajax:
                    return JsonResponse({'status': 'error', 'msg': "Kelgan sana kelajakdagi sana bo'lishi mumkin emas!"}, status=400)
                messages.error(request, "Kelgan sana kelajakdagi sana bo'lishi mumkin emas!")
            elif student.joined_date != new_joined:
                student.joined_date = new_joined
                student.save(update_fields=['joined_date'])
        except Exception:
            pass

    mp, _ = MonthlyPayment.objects.update_or_create(
        student=student,
        month=month,
        defaults={
            'amount_paid': amount_to_save,
            'payment_type': payment_type,
            'payment_date': date.today(),
            'status': new_status,
            'notes': notes,
        },
    )

    try:
        grp_name = student.time_slot.group.name if (student.time_slot and student.time_slot.group) else '-'
        AuditLog.objects.create(
            action_type='payment',
            actor=request.user if getattr(request, 'user', None) and request.user.is_authenticated else None,
            actor_name="Direktor / Admin",
            title=f"To'lov qayd etildi: {student.first_name} {student.last_name} ({amount_to_save} - {mp.get_payment_type_display()})",
            details=f"Guruh: {grp_name}, Oy: {month}, Turi: {mp.get_payment_type_display()}, Holat: {new_status}",
        )
    except Exception:
        pass

    if is_ajax:
        rate_val = student.payment_rate or 0
        due_date_str = ''
        if mp.payment_date and (amount_num > 0 or has_text or is_mahalla_user):
            if relativedelta:
                due_date_obj = mp.payment_date + relativedelta(months=1)
            else:
                yr, mn = mp.payment_date.year, mp.payment_date.month + 1
                if mn > 12:
                    mn, yr = 1, yr + 1
                due_date_obj = date(yr, mn, min(mp.payment_date.day, calendar.monthrange(yr, mn)[1]))
            due_date_str = due_date_obj.strftime('%d.%m.%Y')

        is_paid = (new_status == 'paid') or has_text or is_mahalla_user
        if is_mahalla_user:
            status_display = "Mahalla" if amount_num == 0 else f"To'langan ({amount_to_save})"
            st_code_resp, debt = 'paid', 0
        elif has_text:
            status_display = f"To'langan ({amount_to_save})"
            st_code_resp, debt = 'paid', 0
        elif is_paid:
            status_display = "To'langan"
            st_code_resp, debt = 'paid', 0
        else:
            status_display = "Qarzdor"
            st_code_resp, debt = 'unpaid', rate_val

        return JsonResponse({
            'status': 'success',
            'payment_id': mp.id,
            'amount_paid': mp.amount_paid,
            'amount_paid_num': amount_num,
            'amount_paid_formatted': _fmt_money(amount_num) if amount_num > 0 else (amount_to_save if (has_text or is_mahalla_user) else '0'),
            'payment_type': mp.payment_type,
            'payment_type_display': mp.get_payment_type_display(),
            'payment_date': mp.payment_date.strftime('%d.%m.%Y') if mp.payment_date else '',
            'payment_date_iso': mp.payment_date.strftime('%Y-%m-%d') if mp.payment_date else '',
            'due_date': due_date_str,
            'status_code': st_code_resp,
            'status_display': status_display,
            'notes': mp.notes,
            'rate': rate_val,
            'debt': debt,
            'debt_formatted': _fmt_money(debt) if debt > 0 else '0',
        })

    if amount_num > 0:
        messages.success(request, f"✅ {student}: {_fmt_money(amount_num)} so'm [{mp.get_payment_type_display()}] ({month}) saqlandi!")
    elif is_mahalla_user:
        messages.success(request, f"✅ {student}: Mahalla toifasi ({month}) saqlandi!")
    elif has_text:
        messages.success(request, f"✅ {student}: '{amount_to_save}' ({month}) saqlandi!")
    else:
        messages.warning(request, f"⚠️ {student}: {month} uchun to'lov bekor qilindi (0 so'm).")
    return redirect(request.POST.get('next') or request.META.get('HTTP_REFERER') or 'admin_payments_all')


@admin_required
def admin_payment_delete(request, id):
    """Tanlangan to'lov yozuvini bekor qilish / o'chirish."""
    if not _admin_only(request):
        return redirect('home')
    mp = get_object_or_404(MonthlyPayment, id=id)
    if request.method == 'POST':
        student_name = str(mp.student)
        month = mp.month
        mp.delete()
        messages.success(request, f"{student_name}ning {month} oyi uchun to'lovi o'chirildi.")
    return redirect(request.GET.get('next') or request.POST.get('next') or request.META.get('HTTP_REFERER') or 'admin_payments_all')


@admin_required
def payment_mark(request):
    """Eski payment_mark endpointi bilan moslik (backward compatibility)."""
    return admin_payment_save(request)


@admin_required
def export_payments_excel(request):
    """To'lovlar jadvalini professional Excel fayl (.xlsx) shaklida eksport qilish."""
    if not _admin_only(request):
        return redirect('home')

    req_month, _, _ = normalize_month_str(request.GET.get('month'))
    group_id = request.GET.get('group_id', '').strip()
    status_filter = (request.GET.get('status') or 'all').strip().lower()
    if request.GET.get('is_local') in ('1', 'true', 'True') or status_filter == 'local':
        status_filter = 'local'

    is_grp = group_id.isdigit()
    if is_grp:
        group = get_object_or_404(Group.objects.select_related('teacher'), id=int(group_id))
        students_qs = get_active_students_qs(group=group)
        base_title = f"Guruh To'lovlari — {group.name}"
        safe_name = "".join(c for c in group.name if c.isalnum() or c in (' ', '_', '-')).strip().replace(' ', '_')
        base_filename = f"Toluvlar_{safe_name}"
        group_info = f"Fan: {group.subject_name} | O'qituvchi: {group.teacher or '—'}"
    else:
        students_qs = get_active_students_qs()
        base_title = "Barcha Guruhlar Umumiy To'lovlar Jadvali"
        base_filename = "Humo_Umumiy_Toluvlar"
        group_info = "Barcha guruhlar"

    status_meta = {
        'paid': ("To'liq To'laganlar", f"{base_filename}_Tolaganlar_{req_month}.xlsx"),
        'overdue': ("Kechikkanlar (Muddati O'tganlar)", f"{base_filename}_Kechikkanlar_{req_month}.xlsx"),
        'unpaid': ("Qarzdorlar (Umuman to'lamaganlar)", f"{base_filename}_Qarzdorlar_{req_month}.xlsx"),
        'local': ("Mahalladan Kelganlar", f"{base_filename}_Mahalladan_{req_month}.xlsx"),
    }
    type_label, filename = status_meta.get(status_filter, ("Barcha O'quvchilar", f"{base_filename}_{req_month}.xlsx"))

    title = f"{base_title} ({type_label})"
    subtitle = f"{group_info} | Holat: {type_label} | Davr: {req_month} | Yuklangan: {date.today().strftime('%d.%m.%Y')}"

    filter_for_data = 'all' if status_filter == 'yearly' else status_filter
    rows, stats = _build_payments_data(students_qs, req_month, status_filter=filter_for_data, is_group_specific=is_grp)
    excel_file = generate_payments_excel(rows, title=title, subtitle=subtitle, stats=stats)

    response = HttpResponse(
        excel_file.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


@admin_required
def admin_export_excel(request):
    """Admin: Barcha o'qituvchilar va ularga biriktirilgan guruhlar/o'quvchilar ro'yxatini yuklab berish."""
    if not _admin_only(request):
        return redirect('home')

    curr_month = date.today().strftime('%Y-%m')
    payment_map = dict(MonthlyPayment.objects.filter(month=curr_month).values_list('student_id', 'amount_paid'))

    all_students = list(
        get_active_students_qs().order_by(
            'time_slot__group__teacher__first_name',
            'time_slot__group__teacher__last_name',
            'time_slot__group__name',
            'first_name',
            'last_name',
        )
    )

    wb = openpyxl.Workbook()
    ws_all = wb.active
    ws_all.title = "Barcha O'quvchilar"
    _populate_excel_sheet(ws_all, all_students, payment_map)

    for t in Teacher.objects.order_by('first_name', 'last_name'):
        t_students = [s for s in all_students if s.time_slot and s.time_slot.group and s.time_slot.group.teacher_id == t.id]
        if t_students:
            ws_t = wb.create_sheet(title=f"{t.first_name} {t.last_name}"[:30])
            _populate_excel_sheet(ws_t, t_students, payment_map)

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="HUMO_Oquvchilar_{curr_month}.xlsx"'
    wb.save(response)
    return response


@admin_required
def admin_group_export_excel(request, group_id):
    """Guruh o'quvchilarini Excel fayl qilib yuklab berish."""
    if not _admin_only(request):
        return redirect('home')

    group = get_object_or_404(Group, id=group_id)
    curr_month = date.today().strftime('%Y-%m')
    payment_map = dict(MonthlyPayment.objects.filter(month=curr_month).values_list('student_id', 'amount_paid'))
    students = get_active_students_qs(group=group).order_by('first_name', 'last_name')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = group.name[:28]
    _populate_excel_sheet(ws, students, payment_map)

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="Guruh_{group.name}_{curr_month}.xlsx"'
    wb.save(response)
    return response

